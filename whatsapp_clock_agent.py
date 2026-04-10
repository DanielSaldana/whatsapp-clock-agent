import io
import os
import sqlite3
from datetime import datetime, timezone
from flask import Flask, request, jsonify, Response, render_template_string, send_file
from openpyxl import Workbook
from twilio.rest import Client  # 👈 NUEVO

try:
    import psycopg
    from psycopg.rows import dict_row
except Exception:
    psycopg = None
    dict_row = None

APP_NAME = "WhatsApp Clock Agent"
DB_PATH = os.getenv("DB_PATH", "timeclock.db")
DATABASE_URL = os.getenv("DATABASE_URL", "").strip()
ADMIN_NUMBERS = {n.strip() for n in os.getenv("ADMIN_NUMBERS", "").split(",") if n.strip()}
TIMEZONE_LABEL = os.getenv("TIMEZONE_LABEL", "America/Denver")
ADMIN_TOKEN = os.getenv("ADMIN_TOKEN", "mi-dashboard-2026")

app = Flask(__name__)

# ================================
# 🔑 TWILIO
# ================================
ACCOUNT_SID = os.getenv("ACCOUNT_SID")
AUTH_TOKEN = os.getenv("AUTH_TOKEN")
TWILIO_WHATSAPP_NUMBER = "whatsapp:+19705405717"

# Cliente Twilio lazy: se inicializa solo si se necesita (evita crash al arrancar)
_twilio_client = None
def get_twilio_client():
    global _twilio_client
    if _twilio_client is None:
        if not ACCOUNT_SID or not AUTH_TOKEN:
            raise RuntimeError("ACCOUNT_SID y AUTH_TOKEN no están configurados como variables de entorno.")
        _twilio_client = Client(ACCOUNT_SID, AUTH_TOKEN)
    return _twilio_client


# -----------------------------
# Database
# -----------------------------
def using_postgres() -> bool:
    return DATABASE_URL.startswith("postgres://") or DATABASE_URL.startswith("postgresql://")


def normalized_database_url() -> str:
    if DATABASE_URL.startswith("postgres://"):
        return DATABASE_URL.replace("postgres://", "postgresql://", 1)
    return DATABASE_URL


def get_conn():
    if using_postgres():
        if psycopg is None:
            raise RuntimeError("psycopg is not installed. Add psycopg[binary] to requirements.txt")
        return psycopg.connect(normalized_database_url(), row_factory=dict_row)
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    return conn


def sql_placeholders(sql: str) -> str:
    return sql.replace("?", "%s") if using_postgres() else sql


def db_execute(sql: str, params=(), *, fetchone=False, fetchall=False, commit=False):
    conn = get_conn()
    cur = conn.cursor()
    cur.execute(sql_placeholders(sql), params)
    result = None
    if fetchone:
        result = cur.fetchone()
    elif fetchall:
        result = cur.fetchall()
    if commit:
        conn.commit()
    cur.close()
    conn.close()
    return result


def init_db():
    conn = get_conn()
    cur = conn.cursor()

    if using_postgres():
        cur.execute(
            """
            CREATE TABLE IF NOT EXISTS employees (
                id BIGSERIAL PRIMARY KEY,
                phone TEXT UNIQUE NOT NULL,
                name TEXT NOT NULL
            )
            """
        )
        cur.execute(
            """
            CREATE TABLE IF NOT EXISTS shifts (
                id BIGSERIAL PRIMARY KEY,
                phone TEXT NOT NULL,
                employee_name TEXT,
                date_local TEXT,
                clock_in_utc TEXT,
                clock_out_utc TEXT,
                lunch_minutes INTEGER DEFAULT 0,
                in_lat DOUBLE PRECISION,
                in_lng DOUBLE PRECISION,
                out_lat DOUBLE PRECISION,
                out_lng DOUBLE PRECISION,
                location_description_in TEXT,
                location_description_out TEXT,
                notes TEXT,
                total_work_minutes INTEGER,
                status TEXT DEFAULT 'open'
            )
            """
        )
        cur.execute(
            """
            CREATE TABLE IF NOT EXISTS conversation_state (
                phone TEXT PRIMARY KEY,
                state TEXT,
                temp_value TEXT,
                updated_at_utc TEXT
            )
            """
        )
    else:
        cur.execute(
            """
            CREATE TABLE IF NOT EXISTS employees (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                phone TEXT UNIQUE NOT NULL,
                name TEXT NOT NULL
            )
            """
        )
        cur.execute(
            """
            CREATE TABLE IF NOT EXISTS shifts (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                phone TEXT NOT NULL,
                employee_name TEXT,
                date_local TEXT,
                clock_in_utc TEXT,
                clock_out_utc TEXT,
                lunch_minutes INTEGER DEFAULT 0,
                in_lat REAL,
                in_lng REAL,
                out_lat REAL,
                out_lng REAL,
                location_description_in TEXT,
                location_description_out TEXT,
                notes TEXT,
                total_work_minutes INTEGER,
                status TEXT DEFAULT 'open'
            )
            """
        )
        cur.execute(
            """
            CREATE TABLE IF NOT EXISTS conversation_state (
                phone TEXT PRIMARY KEY,
                state TEXT,
                temp_value TEXT,
                updated_at_utc TEXT
            )
            """
        )

    conn.commit()
    cur.close()
    conn.close()


# -----------------------------
# Utilities
# -----------------------------
def utc_now_iso() -> str:
    return datetime.now(timezone.utc).isoformat()


def local_date_string() -> str:
    return datetime.now().strftime("%Y-%m-%d")


def normalize_text(value: str) -> str:
    return (value or "").strip()


def parse_iso(value: str):
    if not value:
        return None
    return datetime.fromisoformat(value)


def fmt_dt(iso_value: str) -> str:
    dt = parse_iso(iso_value)
    if not dt:
        return "-"
    return dt.strftime("%Y-%m-%d %I:%M %p UTC")


def fmt_minutes(total_minutes: int) -> str:
    total_minutes = int(total_minutes or 0)
    hours = total_minutes // 60
    minutes = total_minutes % 60
    return f"{hours}h {minutes}m"


def admin_authorized(req) -> bool:
    token = req.args.get("token") or req.form.get("token") or req.headers.get("X-Admin-Token", "")
    return bool(ADMIN_TOKEN) and token == ADMIN_TOKEN


def parse_time_token(raw: str):
    raw = normalize_text(raw).lower()
    if not raw:
        return None
    normalized = (
        raw.replace("a.m.", "am")
        .replace("p.m.", "pm")
        .replace("a.m", "am")
        .replace("p.m", "pm")
        .replace("am", " am")
        .replace("pm", " pm")
    )
    normalized = " ".join(normalized.split())
    patterns = ["%I:%M %p", "%I %p", "%H:%M", "%H%M"]
    for pattern in patterns:
        try:
            return datetime.strptime(normalized, pattern)
        except ValueError:
            pass
    return None


def time_to_today(raw: str):
    parsed = parse_time_token(raw)
    if not parsed:
        return None
    now = datetime.now()
    return datetime(now.year, now.month, now.day, parsed.hour, parsed.minute)


def parse_turno_message(text: str):
    t = normalize_text(text).lower()
    starts = ["turno ", "turn ", "trno ", "truno ", "tuno "]
    if not any(t.startswith(s) for s in starts):
        return None

    payload = t.split(" ", 1)[1].strip()
    replacements = {
        "minutos": "min",
        "mins": "min",
        "minuto": "min",
        "lunch": "lonche",
        "lonch": "lonche",
        "lonce": "lonche",
        "lonhce": "lonche",
        "lonchee": "lonche",
        "lounche": "lonche",
    }
    for old, new in replacements.items():
        payload = payload.replace(old, new)
    payload = payload.replace("–", "-").replace("—", "-")

    if "lonche" not in payload:
        return None

    before, site = payload.split("lonche", 1)
    site = normalize_text(site).title()
    parts = before.split()

    try:
        if len(parts) >= 2 and "-" in parts[0]:
            start_raw, end_raw = parts[0].split("-", 1)
            lunch_raw = parts[1]
        elif len(parts) >= 5:
            start_raw = f"{parts[0]} {parts[1]}"
            end_raw = f"{parts[2]} {parts[3]}"
            lunch_raw = parts[4]
        elif len(parts) >= 3:
            start_raw = parts[0]
            end_raw = parts[1]
            lunch_raw = parts[2]
        else:
            return None

        lunch_minutes = int("".join(ch for ch in lunch_raw if ch.isdigit()))
        start_dt = time_to_today(start_raw)
        end_dt = time_to_today(end_raw)
        if not start_dt or not end_dt:
            return None
        if end_dt <= start_dt:
            return None
        return start_dt, end_dt, lunch_minutes, site
    except Exception:
        return None


# -----------------------------
# Persistence helpers
# -----------------------------
def set_state(phone: str, state: str, temp_value: str = ""):
    if using_postgres():
        db_execute(
            """
            INSERT INTO conversation_state(phone, state, temp_value, updated_at_utc)
            VALUES(%s, %s, %s, %s)
            ON CONFLICT(phone) DO UPDATE SET
                state = EXCLUDED.state,
                temp_value = EXCLUDED.temp_value,
                updated_at_utc = EXCLUDED.updated_at_utc
            """,
            (phone, state, temp_value, utc_now_iso()),
            commit=True,
        )
    else:
        db_execute(
            """
            INSERT OR REPLACE INTO conversation_state(phone, state, temp_value, updated_at_utc)
            VALUES(?, ?, ?, ?)
            """,
            (phone, state, temp_value, utc_now_iso()),
            commit=True,
        )


def get_state(phone: str):
    return db_execute("SELECT * FROM conversation_state WHERE phone = ?", (phone,), fetchone=True)


def clear_state(phone: str):
    db_execute("DELETE FROM conversation_state WHERE phone = ?", (phone,), commit=True)


def get_employee_name(phone: str):
    row = db_execute("SELECT name FROM employees WHERE phone = ?", (phone,), fetchone=True)
    return row["name"] if row else None


def set_employee_name(phone: str, name: str):
    if using_postgres():
        db_execute(
            """
            INSERT INTO employees(phone, name)
            VALUES(%s, %s)
            ON CONFLICT(phone) DO UPDATE SET name = EXCLUDED.name
            """,
            (phone, name),
            commit=True,
        )
    else:
        db_execute(
            """
            INSERT OR REPLACE INTO employees(phone, name)
            VALUES(?, ?)
            """,
            (phone, name),
            commit=True,
        )


def get_open_shift(phone: str):
    return db_execute(
        "SELECT * FROM shifts WHERE phone = ? AND status = 'open' ORDER BY id DESC LIMIT 1",
        (phone,),
        fetchone=True,
    )


def create_shift(phone: str, employee_name: str, loc_description: str = "", lat=None, lng=None):
    db_execute(
        """
        INSERT INTO shifts(
            phone, employee_name, date_local, clock_in_utc,
            in_lat, in_lng, location_description_in, status
        )
        VALUES(?, ?, ?, ?, ?, ?, ?, 'open')
        """,
        (phone, employee_name, local_date_string(), utc_now_iso(), lat, lng, loc_description),
        commit=True,
    )


def create_shift_manual(phone: str, employee_name: str, clock_in_dt: datetime, clock_out_dt: datetime, lunch_minutes: int, site: str):
    total_minutes = max(int((clock_out_dt - clock_in_dt).total_seconds() // 60) - int(lunch_minutes), 0)
    db_execute(
        """
        INSERT INTO shifts(
            phone, employee_name, date_local, clock_in_utc, clock_out_utc,
            lunch_minutes, location_description_in, location_description_out,
            total_work_minutes, status
        )
        VALUES(?, ?, ?, ?, ?, ?, ?, ?, ?, 'closed')
        """,
        (
            phone,
            employee_name,
            clock_in_dt.strftime("%Y-%m-%d"),
            clock_in_dt.replace(tzinfo=timezone.utc).isoformat(),
            clock_out_dt.replace(tzinfo=timezone.utc).isoformat(),
            int(lunch_minutes),
            site,
            site,
            total_minutes,
        ),
        commit=True,
    )
    return total_minutes


def close_shift(phone: str, lunch_minutes: int = 0, notes: str = "", loc_description: str = "", lat=None, lng=None):
    open_shift = get_open_shift(phone)
    if not open_shift:
        return None
    clock_in = datetime.fromisoformat(open_shift["clock_in_utc"])
    clock_out = datetime.fromisoformat(utc_now_iso())
    total_minutes = max(int((clock_out - clock_in).total_seconds() // 60) - int(lunch_minutes), 0)
    db_execute(
        """
        UPDATE shifts
        SET clock_out_utc = ?,
            lunch_minutes = ?,
            out_lat = ?,
            out_lng = ?,
            location_description_out = ?,
            notes = ?,
            total_work_minutes = ?,
            status = 'closed'
        WHERE id = ?
        """,
        (
            clock_out.isoformat(),
            int(lunch_minutes),
            lat,
            lng,
            loc_description,
            notes,
            total_minutes,
            open_shift["id"],
        ),
        commit=True,
    )
    return total_minutes


def latest_closed_shifts(limit: int = 20):
    return db_execute(
        """
        SELECT employee_name, phone, date_local, lunch_minutes, total_work_minutes,
               location_description_in, location_description_out,
               clock_in_utc, clock_out_utc, notes
        FROM shifts
        WHERE status = 'closed'
        ORDER BY id DESC
        LIMIT ?
        """,
        (limit,),
        fetchall=True,
    )


def fetch_dashboard_shifts(employee: str = "", date_from: str = "", date_to: str = ""):
    sql = """
        SELECT id, employee_name, phone, date_local, clock_in_utc, clock_out_utc,
               lunch_minutes, location_description_in, location_description_out,
               in_lat, in_lng, out_lat, out_lng, notes, total_work_minutes, status
        FROM shifts
        WHERE 1=1
    """
    params = []
    if employee:
        sql += " AND (employee_name LIKE ? OR phone LIKE ?)"
        like = f"%{employee}%"
        params.extend([like, like])
    if date_from:
        sql += " AND date_local >= ?"
        params.append(date_from)
    if date_to:
        sql += " AND date_local <= ?"
        params.append(date_to)
    sql += " ORDER BY date_local DESC, id DESC"
    return db_execute(sql, tuple(params), fetchall=True)


def build_dashboard_summary(rows):
    total_minutes = sum(int(row["total_work_minutes"] or 0) for row in rows if row["status"] == "closed")
    unique_employees = len({(row["employee_name"] or row["phone"]) for row in rows})
    open_shifts = sum(1 for row in rows if row["status"] == "open")
    closed_shifts = sum(1 for row in rows if row["status"] == "closed")
    return {
        "total_minutes": total_minutes,
        "unique_employees": unique_employees,
        "open_shifts": open_shifts,
        "closed_shifts": closed_shifts,
    }


def build_excel(rows):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Shifts"
    headers = [
        "Employee", "Phone", "Date", "Clock In UTC", "Clock Out UTC", "Lunch Minutes",
        "Total Worked Minutes", "Total Worked", "Status", "Clock In Site", "Clock Out Site",
        "In Latitude", "In Longitude", "Out Latitude", "Out Longitude", "Notes",
    ]
    sheet.append(headers)
    for row in rows:
        sheet.append([
            row["employee_name"], row["phone"], row["date_local"], row["clock_in_utc"], row["clock_out_utc"],
            row["lunch_minutes"], row["total_work_minutes"], fmt_minutes(row["total_work_minutes"] or 0),
            row["status"], row["location_description_in"], row["location_description_out"],
            row["in_lat"], row["in_lng"], row["out_lat"], row["out_lng"], row["notes"],
        ])
    for column in sheet.columns:
        max_length = 0
        letter = column[0].column_letter
        for cell in column:
            max_length = max(max_length, len("" if cell.value is None else str(cell.value)))
        sheet.column_dimensions[letter].width = min(max_length + 2, 28)
    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


# -----------------------------
# Messaging
# -----------------------------
def help_text() -> str:
    return (
        "*Clock Agent Commands*\n\n"
        "start - register your name\n"
        "in - clock in\n"
        "out - clock out\n"
        "turno 8:00am 5:30pm 30min lonche rancho - save full shift in one message\n"
        "lunch 30 - set lunch minutes\n"
        "site Warehouse A - add work location description\n"
        "note Worked on delivery route - add notes before clock out\n"
        "status - view current open shift\n"
        "help - show commands"
    )


def parse_whatsapp_location(form):
    lat = form.get("Latitude") or form.get("WaLatitude")
    lng = form.get("Longitude") or form.get("WaLongitude")
    try:
        if lat is not None and lng is not None:
            return float(lat), float(lng)
    except ValueError:
        return None, None
    return None, None


def incoming_text(form) -> str:
    return normalize_text(form.get("Body", ""))


def from_number(form) -> str:
    return normalize_text(form.get("From", ""))


def is_location_message(form) -> bool:
    lat, lng = parse_whatsapp_location(form)
    return lat is not None and lng is not None


def save_location_to_open_shift(phone: str, lat: float, lng: float):
    shift = get_open_shift(phone)
    if not shift:
        return False, "No open shift found. Send *in* first."
    if shift["in_lat"] is None or shift["in_lng"] is None:
        db_execute("UPDATE shifts SET in_lat = ?, in_lng = ? WHERE id = ?", (lat, lng, shift["id"]), commit=True)
        return True, "Location saved for *clock in*."
    db_execute("UPDATE shifts SET out_lat = ?, out_lng = ? WHERE id = ?", (lat, lng, shift["id"]), commit=True)
    return True, "Location saved for current shift."


def handle_stateful_reply(phone: str, text: str):
    state = get_state(phone)
    if not state:
        return None
    current = state["state"]

    if current == "awaiting_name":
        name = text.strip()
        if not name:
            return "Send your full name to complete registration."
        set_employee_name(phone, name)
        clear_state(phone)
        return f"Perfecto. Te registraste como *{name}*. Envia tu turno Ejemplo ( Turno 8:00am 4:30pm 30min lonche Rastrillar."

    if current == "awaiting_in_description":
        name = get_employee_name(phone) or "Employee"
        create_shift(phone, name, loc_description=text)
        clear_state(phone)
        return (
            f"✅ Clock in saved for *{name}*.\n"
            f"Site: {text or 'Not provided'}\n"
            "Now send your WhatsApp location if you want GPS attached."
        )

    if current == "awaiting_out_lunch":
        try:
            lunch_minutes = int(text)
            if lunch_minutes < 0 or lunch_minutes > 240:
                return "Send lunch as minutes only, for example: 30"
        except ValueError:
            return "Send lunch as minutes only, for example: 30"
        set_state(phone, "awaiting_out_description", str(lunch_minutes))
        return "Got it. Now send the location description for clock out. Example: *Warehouse B*"

    if current == "awaiting_out_description":
        try:
            lunch_minutes = int(state["temp_value"] or "0")
        except ValueError:
            lunch_minutes = 0
        total = close_shift(phone, lunch_minutes=lunch_minutes, loc_description=text)
        clear_state(phone)
        if total is None:
            return "No open shift found. Send *in* first."
        return (
            f"✅ Clock out saved.\n"
            f"Lunch: {lunch_minutes} min\n"
            f"Site: {text or 'Not provided'}\n"
            f"Worked: *{fmt_minutes(total)}*"
        )

    return None


def handle_command(phone: str, text: str):
    turno = parse_turno_message(text)
    if turno:
        name = get_employee_name(phone)
        if not name:
            set_state(phone, "awaiting_name")
            return "First register your name. Send your full name."
        if get_open_shift(phone):
            return "You already have an open shift. Send *out* first."
        start_dt, end_dt, lunch_minutes, site = turno
        total = create_shift_manual(phone, name, start_dt, end_dt, lunch_minutes, site)
        return (
            "✅ Turno guardado en un mensaje.\n"
            f"Entrada: *{start_dt.strftime('%I:%M %p')}*\n"
            f"Salida: *{end_dt.strftime('%I:%M %p')}*\n"
            f"Lonche: {lunch_minutes} min\n"
            f"Lugar: {site or 'No especificado'}\n"
            f"Total trabajado: *{fmt_minutes(total)}*"
        )

    lower = text.lower().strip()
    name = get_employee_name(phone)

    if lower in {"help", "menu", "ayuda"}:
        return help_text()
    if lower.startswith("start"):
        name = text.replace("start", "").strip()
        if not name:
            set_state(phone, "awaiting_name")
            return "Envia tu Nombre y apellido"
        set_employee_name(phone, name)
        create_shift(phone, name, loc_description="Auto")
        return f"✅ Registrado como {name} y turno iniciado"
    if lower == "status":
        shift = get_open_shift(phone)
        if not shift:
            return "You do not have an open shift. Send *in* to clock in."
        return (
            f"Open shift found for *{shift['employee_name'] or 'Employee'}*.\n"
            f"Date: {shift['date_local']}\n"
            f"Clock in: {shift['clock_in_utc']} UTC\n"
            f"Site: {shift['location_description_in'] or 'Not provided'}"
        )
    if lower == "in":
        if not name:
            set_state(phone, "awaiting_name")
            return "First register your name. Send your full name."
        if get_open_shift(phone):
            return "You already have an open shift. Send *status* or *out*."
        set_state(phone, "awaiting_in_description")
        return "Send the work location description for clock in. Example: *Job Site 4*"
    if lower == "out":
        if not get_open_shift(phone):
            return "No open shift found. Send *in* first."
        set_state(phone, "awaiting_out_lunch")
        return "Send lunch minutes only. Example: *30*"
    if lower.startswith("lunch "):
        try:
            minutes = int(lower.split(" ", 1)[1].strip())
        except ValueError:
            return "Use: *lunch 30*"
        shift = get_open_shift(phone)
        if not shift:
            return "No open shift found. Send *in* first."
        db_execute("UPDATE shifts SET lunch_minutes = ? WHERE id = ?", (minutes, shift["id"]), commit=True)
        return f"Lunch updated to *{minutes}* minutes for your current shift."
    if lower.startswith("site "):
        site_text = text[5:].strip()
        shift = get_open_shift(phone)
        if not shift:
            return "No open shift found. Send *in* first."
        db_execute("UPDATE shifts SET location_description_in = ? WHERE id = ?", (site_text, shift["id"]), commit=True)
        return f"Site saved: *{site_text}*"
    if lower.startswith("note "):
        note_text = text[5:].strip()
        shift = get_open_shift(phone)
        if not shift:
            return "No open shift found. Send *in* first."
        db_execute("UPDATE shifts SET notes = ? WHERE id = ?", (note_text, shift["id"]), commit=True)
        return "Note saved to your current shift."
    if lower == "report":
        if phone not in ADMIN_NUMBERS:
            return "This command is only available for admin numbers."
        rows = latest_closed_shifts(limit=10)
        if not rows:
            return "No closed shifts yet."
        lines = ["*Latest closed shifts*\n"]
        for row in rows:
            lines.append(f"• {row['employee_name'] or row['phone']} | {row['date_local']} | {fmt_minutes(row['total_work_minutes'] or 0)}")
        return "\n".join(lines)
    return f"No entiendo eso.\n\n{help_text()}"


# -----------------------------
# Web routes
# -----------------------------
@app.route("/")
def home():
    return jsonify({
        "app": APP_NAME,
        "ok": True,
        "database": "postgres" if using_postgres() else "sqlite",
        "message": "Use /health, /dashboard?token=YOUR_TOKEN or /export.xlsx?token=YOUR_TOKEN",
    })


@app.route("/health", methods=["GET"])
def health():
    return jsonify({
        "ok": True,
        "app": APP_NAME,
        "timezone": TIMEZONE_LABEL,
        "database": "postgres" if using_postgres() else "sqlite",
    })


@app.route("/dashboard", methods=["GET"])
def dashboard():
    if not admin_authorized(request):
        return Response("Unauthorized", status=401)

    employee = normalize_text(request.args.get("employee", ""))
    date_from = normalize_text(request.args.get("date_from", ""))
    date_to = normalize_text(request.args.get("date_to", ""))
    rows = fetch_dashboard_shifts(employee=employee, date_from=date_from, date_to=date_to)
    summary = build_dashboard_summary(rows)
    html = """
<!DOCTYPE html>
<html lang="es">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>Clock Agent Dashboard</title>
<style>
* { box-sizing: border-box; margin: 0; padding: 0; }

body {
  font-family: 'Inter', sans-serif;
  background: #0b0f1a;
  color: #e6edf3;
}

header {
  background: rgba(20, 25, 40, 0.7);
  backdrop-filter: blur(12px);
  border-bottom: 1px solid rgba(255,255,255,0.05);
  padding: 18px 28px;
}

header h1 {
  font-size: 1.3rem;
  font-weight: 600;
  color: #fff;
}

.container {
  max-width: 1200px;
  margin: 30px auto;
  padding: 0 20px;
}

.cards {
  display: flex;
  gap: 16px;
  flex-wrap: wrap;
  margin-bottom: 30px;
}

.card {
  background: rgba(20, 25, 40, 0.8);
  border: 1px solid rgba(255,255,255,0.05);
  border-radius: 14px;
  padding: 24px;
  flex: 1;
  min-width: 180px;
  backdrop-filter: blur(10px);
  transition: 0.2s;
}

.card:hover {
  transform: translateY(-3px);
  box-shadow: 0 10px 30px rgba(0,0,0,0.4);
}

.card .num {
  font-size: 2.2rem;
  font-weight: bold;
  color: #4da3ff;
}

.card .lbl {
  font-size: 0.85rem;
  color: #8b949e;
  margin-top: 6px;
}

.filters {
  background: rgba(20, 25, 40, 0.8);
  border: 1px solid rgba(255,255,255,0.05);
  border-radius: 12px;
  padding: 18px;
  margin-bottom: 24px;
  display: flex;
  gap: 12px;
  flex-wrap: wrap;
}

.filters input {
  background: #0b0f1a;
  border: 1px solid rgba(255,255,255,0.1);
  color: white;
  padding: 8px 10px;
  border-radius: 6px;
}

.filters button {
  background: #4da3ff;
  border: none;
  color: white;
  padding: 8px 18px;
  border-radius: 6px;
  cursor: pointer;
}

.export-btn {
  background: #22c55e;
  padding: 8px 16px;
  border-radius: 6px;
  color: white;
  text-decoration: none;
}

table {
  width: 100%;
  border-collapse: collapse;
  background: rgba(20, 25, 40, 0.8);
  border-radius: 12px;
  overflow: hidden;
}

th {
  background: #111827;
  color: #9ca3af;
  padding: 12px;
  font-size: 0.8rem;
  text-transform: uppercase;
}

td {
  padding: 12px;
  border-top: 1px solid rgba(255,255,255,0.05);
}

tr:hover td {
  background: rgba(255,255,255,0.03);
}

.badge-open {
  background: rgba(250,204,21,0.15);
  color: #facc15;
  padding: 3px 10px;
  border-radius: 10px;
}

.badge-closed {
  background: rgba(34,197,94,0.15);
  color: #22c55e;
  padding: 3px 10px;
  border-radius: 10px;
}

.empty {
  text-align: center;
  padding: 50px;
  color: #6b7280;
}
</style>
</head>
<body>
<header><h1>⏱ Clock Agent Dashboard</h1></header>
<div class="container">
    <div style="margin-bottom:20px;">
  <form method="post" action="/create-employee" style="display:flex;gap:10px;flex-wrap:wrap;">
    
    <input type="hidden" name="token" value="{{ token }}">

    <input type="text" name="name" placeholder="Nombre del empleado"
      style="padding:10px;border-radius:6px;border:none;background:#111827;color:white;">

    <input type="text" name="phone" placeholder="whatsapp:+1..."
      style="padding:10px;border-radius:6px;border:none;background:#111827;color:white;">

    <button type="submit"
      style="background:#4da3ff;border:none;padding:10px 16px;border-radius:6px;color:white;cursor:pointer;">
      ➕ Crear empleado
    </button>

  </form>
</div>
  <div class="cards">
    <div class="card"><div class="num">{{ summary.unique_employees }}</div><div class="lbl">Empleados</div></div>
    <div class="card"><div class="num">{{ summary.open_shifts }}</div><div class="lbl">Turnos Abiertos</div></div>
    <div class="card"><div class="num">{{ summary.closed_shifts }}</div><div class="lbl">Turnos Cerrados</div></div>
    <div class="card"><div class="num">{{ worked_hours }}</div><div class="lbl">Horas Trabajadas</div></div>
  </div>

  <div class="filters">
    <form method="get" style="display:flex;gap:12px;flex-wrap:wrap;align-items:flex-end;width:100%">
      <input type="hidden" name="token" value="{{ token }}">
      <div><label>Empleado</label><input type="text" name="employee" value="{{ employee }}" placeholder="Nombre o teléfono"></div>
      <div><label>Desde</label><input type="date" name="date_from" value="{{ date_from }}"></div>
      <div><label>Hasta</label><input type="date" name="date_to" value="{{ date_to }}"></div>
      <button type="submit">🔍 Filtrar</button>
      <a class="export-btn" href="/export.xlsx?token={{ token }}&employee={{ employee }}&date_from={{ date_from }}&date_to={{ date_to }}">⬇ Excel</a>
    </form>
  </div>

  {% if rows %}
  <table>
    <thead>
      <tr>
        <th>Empleado</th><th>Teléfono</th><th>Fecha</th>
        <th>Entrada</th><th>Salida</th><th>Lonche</th>
        <th>Trabajado</th><th>Lugar Entrada</th><th>Lugar Salida</th>
        <th>Notas</th><th>Estado</th>
      </tr>
    </thead>
    <tbody>
    {% for row in rows %}
      <tr>
        <td>{{ row['employee_name'] or '-' }}</td>
        <td>{{ row['phone'] }}</td>
        <td>{{ row['date_local'] or '-' }}</td>
        <td>{{ fmt_dt(row['clock_in_utc']) }}</td>
        <td>{{ fmt_dt(row['clock_out_utc']) }}</td>
        <td>{{ row['lunch_minutes'] or 0 }} min</td>
        <td>{{ fmt_minutes(row['total_work_minutes'] or 0) }}</td>
        <td>{{ row['location_description_in'] or '-' }}</td>
        <td>{{ row['location_description_out'] or '-' }}</td>
        <td>{{ row['notes'] or '-' }}</td>
        <td>
          {% if row['status'] == 'open' %}
            <span class="badge badge-open">Abierto</span>
          {% else %}
            <span class="badge badge-closed">Cerrado</span>
          {% endif %}
        </td>
      </tr>
    {% endfor %}
    </tbody>
  </table>
  {% else %}
  <div class="empty">📭 No se encontraron registros con los filtros actuales.</div>
  {% endif %}
</div>
</body>
</html>
"""

    return render_template_string(
        html,
        rows=rows,
        summary=summary,
        fmt_minutes=fmt_minutes,
        fmt_dt=fmt_dt,
        worked_hours=fmt_minutes(summary["total_minutes"]),
        token=request.args.get("token", ""),
        employee=employee,
        date_from=date_from,
        date_to=date_to,
    )


@app.route("/export.xlsx", methods=["GET"])
def export_xlsx():
    if not admin_authorized(request):
        return Response("Unauthorized", status=401)
    employee = normalize_text(request.args.get("employee", ""))
    date_from = normalize_text(request.args.get("date_from", ""))
    date_to = normalize_text(request.args.get("date_to", ""))
    rows = fetch_dashboard_shifts(employee=employee, date_from=date_from, date_to=date_to)
    output = build_excel(rows)
    filename = f"clock-report-{datetime.now().strftime('%Y-%m-%d-%H%M')}.xlsx"
    return send_file(
        output,
        as_attachment=True,
        download_name=filename,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
@app.route("/create-employee", methods=["POST"])
def create_employee():
    if not admin_authorized(request):
        return Response("Unauthorized", status=401)

    name = request.form.get("name", "").strip()
    phone = request.form.get("phone", "").strip()

    if not name or not phone:
        return "Missing data", 400

    set_employee_name(phone, name)

    return "OK"

@app.route("/reset-db")
def reset_db():
    token = request.args.get("token")
    if token != os.getenv("ADMIN_TOKEN"):
        return "Unauthorized", 403
    db_execute("DELETE FROM shifts", commit=True)
    db_execute("DELETE FROM employees", commit=True)
    db_execute("DELETE FROM conversation_state", commit=True)
    return "✅ Database reset successful"

@app.route("/reset-shifts")
def reset_shifts():
    token = request.args.get("token")
    if token != os.getenv("ADMIN_TOKEN"):
        return "Unauthorized", 403

    db_execute("DELETE FROM shifts", commit=True)
    db_execute("DELETE FROM conversation_state", commit=True)

    return "✅ Shifts reset only"

# ================================
# 🚀 WEBHOOK (ARREGLADO)
# ================================
from twilio.twiml.messaging_response import MessagingResponse
@app.route("/whatsapp", methods=["POST"])
def whatsapp_webhook():
    print("🔥 HIT /whatsapp")

    form = request.form
    print("FORM:", form)

    phone = from_number(form)
    text = incoming_text(form)

    resp = MessagingResponse()

    # 📍 LOCATION
    if is_location_message(form):
        lat, lng = parse_whatsapp_location(form)
        _, msg = save_location_to_open_shift(phone, lat, lng)
        resp.message(msg)
        return str(resp)

    # 💬 TEXT
    if text:
        state_reply = handle_stateful_reply(phone, text)
        if state_reply:
            resp.message(state_reply)
            return str(resp)

        reply = handle_command(phone, text)
        resp.message(reply)
        return str(resp)

    # 🧠 DEFAULT
    resp.message("Send *help* to see available commands.")
    return str(resp)

init_db()

if __name__ == "__main__":
    port = int(os.getenv("PORT", "5000"))
    app.run(host="0.0.0.0", port=port, debug=False)
